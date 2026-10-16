"""Konvertiert Blutwerte.xlsx in welt-getaggtes Seed-JSON. Nur lesen, nie überschreiben."""
import argparse
import json
import os
import re
from datetime import datetime

SCHEMA_VERSION = 1

def seed_dateiname(welt: str, when: str | None = None) -> str:
    if when is None:
        when = datetime.now().strftime("%Y-%m-%d_%H%M")
    return f"seed_{welt}_{when}.json"

def baue_huelle(welt: str, when_iso: str | None = None) -> dict:
    if when_iso is None:
        when_iso = datetime.now().isoformat(timespec="seconds")
    return {
        "welt": welt,
        "schema_version": SCHEMA_VERSION,
        "exportiert_am": when_iso,
        "profil": None,
        "tageseintraege": {},
        "ereignisse": [],
    }

def parse_marker(text):
    """Interpretiert einen Medikamenten-Marker aus Spalte F/H.
    MJ: eindeutig in mg (vor '12/x' = 10 mg, ab '12/x' = 12,5 mg).
    OZ/WG/Ozempic: nicht rekonstruierbar -> als Rohtext erhalten.
    """
    if not text or not str(text).strip():
        return None
    roh = str(text).strip()
    oben = roh.upper()

    if oben.startswith("MJ"):
        m = re.search(r"(\d+)\s*/\s*\d+", roh)
        dosis = 10
        if m and int(m.group(1)) >= 12:
            dosis = 12.5
        return {"praeparat": "MJ", "dosis_mg": dosis, "pen_anwendung": None, "historie_rohtext": None}

    if oben.startswith("OZ") or oben.startswith("OZEMPIC"):
        return {"praeparat": "OZ", "dosis_mg": None, "pen_anwendung": None, "historie_rohtext": roh}

    if oben.startswith("WG"):
        return {"praeparat": "WG", "dosis_mg": None, "pen_anwendung": None, "historie_rohtext": roh}

    return None

def zeile_zu_eintrag(datum, gewicht, marker):
    """Baut (datum_iso, eintrag) aus einer Tabellenzeile. None, wenn Datum oder Gewicht fehlt."""
    if datum is None or gewicht is None:
        return None
    try:
        gewicht_f = float(gewicht)
    except (TypeError, ValueError):
        return None
    datum_iso = datum.strftime("%Y-%m-%d") if hasattr(datum, "strftime") else str(datum)
    eintrag = {"gewicht_kg": gewicht_f, "spritze": parse_marker(marker)}
    return datum_iso, eintrag

def schreibe_seed(huelle: dict, out_dir: str, when: str | None = None) -> str:
    """Schreibt die Seed-JSON. Überschreibt NIE eine existierende Datei (FileExistsError)."""
    name = seed_dateiname(huelle["welt"], when=when)
    pfad = os.path.join(out_dir, name)
    if os.path.exists(pfad):
        raise FileExistsError(pfad)
    with open(pfad, "x", encoding="utf-8") as f:
        json.dump(huelle, f, ensure_ascii=False, indent=2)
    return pfad

def lies_workbook(xlsx_pfad: str) -> dict:
    """Liest die 3 Gewichts-Tabs und baut die Tageseinträge. Nur lesen (read_only)."""
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_pfad, read_only=True, data_only=True)
    eintraege = {}
    for tab in ("Gewicht 2024", "Gewicht 2025", "Gewicht 2026"):
        if tab not in wb.sheetnames:
            continue
        ws = wb[tab]
        ist_2024 = tab.endswith("2024")
        for row in ws.iter_rows(min_row=2):
            if ist_2024:
                datum = row[0].value if len(row) > 0 else None
                gewicht = row[1].value if len(row) > 1 else None
                marker = None
            else:
                datum = row[2].value if len(row) > 2 else None
                gewicht = row[3].value if len(row) > 3 else None
                f = row[5].value if len(row) > 5 else None
                h = row[7].value if len(row) > 7 else None
                marker = f if isinstance(f, str) else (h if isinstance(h, str) else None)
            res = zeile_zu_eintrag(datum, gewicht, marker)
            if res is not None:
                datum_iso, eintrag = res
                eintraege[datum_iso] = eintrag
    wb.close()
    return eintraege

def main(argv=None):
    p = argparse.ArgumentParser(description="Blutwerte.xlsx -> Seed-JSON")
    p.add_argument("xlsx", help="Pfad zur Blutwerte.xlsx (nur lesen)")
    p.add_argument("--welt", required=True, choices=["test", "produktiv"])
    p.add_argument("--out-dir", default=".")
    args = p.parse_args(argv)
    huelle = baue_huelle(args.welt)
    huelle["tageseintraege"] = lies_workbook(args.xlsx)
    pfad = schreibe_seed(huelle, out_dir=args.out_dir)
    print(f"geschrieben: {pfad}  ({len(huelle['tageseintraege'])} Tage)")

if __name__ == "__main__":
    main()
